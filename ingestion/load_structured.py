import re
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

RAW_XLSX = Path(__file__).resolve().parent.parent / "data" / "raw" / "Hospital_Master_Data.xlsx"
OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "processed" / "structured"

CURRENCY_COLS = {"fee_inr", "price_inr"}
INT_COLS = {"total_beds", "icu_beds", "operation_theatres", "experience_years"}

SHEET_SKIP = {"Index"}

SHEET_CONFIG = {
    "Doctors": (
        "doctors.csv",
        {
            "Doctor ID": "doctor_id",
            "Doctor Name": "doctor_name",
            "Department": "department",
            "Designation": "designation",
            "Qualification": "qualification",
            "Experience (Years)": "experience_years",
            "Consultation Days": "consultation_days",
            "Consultation Timing": "consultation_timing",
            "Fee (INR)": "fee_inr",
            "Room / Location": "room_location",
            "Contact Ext": "contact_ext",
        }
    ),
    "Departments & Facilities": (
        "departments_facilities.csv",
        {
            "Department": "department",
            "Floor / Block": "floor_block",
            "Total Beds": "total_beds",
            "ICU Beds": "icu_beds",
            "Operation Theatres": "operation_theatres",
            "Head of Department": "head_of_department",
            "24x7 Emergency": "emergency_24x7",
            "Contact Ext": "contact_ext",
        }
    ),
    "Ambulance Fleet": (
        "ambulance_fleet.csv",
        {
            "Ambulance ID": "ambulance_id",
            "Vehicle Type": "vehicle_type",
            "Driver Name": "driver_name",
            "Contact Ext": "contact_ext",
        }
    ),
    "Lab Tests & Diagnostics": (
        "lab_tests.csv",
        {
            "Test ID": "test_id",
            "Test Name": "test_name",
            "Department": "department",
            "Cost (INR)": "cost_inr",
            "Sample Type": "sample_type",
            "Turnaround Time": "turnaround_time",
        }
    ),
    "Health Packages": (
        "health_packages.csv",
        {
            "Package ID": "package_id",
            "Package Name": "package_name",
            "Tests Included": "tests_included",
            "Cost (INR)": "cost_inr",
            "Validity (Days)": "validity_days",
        }
    ),
    "Pharmacy": (
        "pharmacy.csv",
        {
            "Medicine ID": "medicine_id",
            "Medicine Name": "medicine_name",
            "Category": "category",
            "Price (INR)": "price_inr",
            "Stock Quantity": "stock_quantity",
            "Expiry Date": "expiry_date",
        }
    ),
    "Insurance & TPA": (
        "insurance_tpa.csv",
        {
            "Insurance ID": "insurance_id",
            "Insurance Name": "insurance_name",
            "TPA Name": "tpa_name",
            "Contact Ext": "contact_ext",
        }
    ),
    "FAQs":(
        "faqs.csv",
        {
            "Question": "question",
            "Answer": "answer",
        }
    )
}

def clean_currency(series: pd.Series)-> pd.Series:
    return (
        series.astype(str)
        .str.replace("₹", "")
        .str.replace(",", "")
        .str.strip()
        .replace({"": None, "nan": None})
        .astype("Int64")
    )

def clean_column_names(df: pd.DataFrame, rename_map: dict) -> pd.DataFrame:
    df = df.rename(columns=rename_map)
    df = df.dropna(how="all")
    for col in df.columns:
        if df[col].dtype == "object":
            df[col] = df[col].astype(str).str.strip()
    for col in CURRENCY_COLS & set(df.columns):
        df[col] = clean_currency(df[col])
    for col in INT_COLS & set(df.columns):
        df[col] = pd.to_numeric(df[col], errors="coerce").astype("Int64")
    return df.reset_index(drop=True)

def load_structured_data():
    wb = load_workbook(RAW_XLSX, data_only=True)

    summary = []
    for sheet_name in wb.sheetnames:
        if sheet_name in SHEET_SKIP:
            continue
        if sheet_name not in SHEET_CONFIG:
            print(f"Warning: No configuration found for sheet '{sheet_name}'. Skipping.")
            continue
        out_name, rename_map = SHEET_CONFIG[sheet_name]
        df = pd.read_excel(RAW_XLSX, sheet_name=sheet_name, engine="openpyxl")
        df = clean_column_names(df, rename_map)

        out_path = OUT_DIR / out_name
        df.to_csv(out_path, index=False)
        summary.append((sheet_name, out_name, len(df), list(df.columns)))

    print(f"\nConverted {len(summary)} sheets from {RAW_XLSX.name} -> {OUT_DIR}/\n")


if __name__ == "__main__":
    load_structured_data()